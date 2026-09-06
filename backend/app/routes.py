import io
import time
from typing import Optional, Dict, Any, List
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Body
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import SalaryHead
from app.schemas import (
    EmployeeResponse,
    SalaryHeadResponse,
    PayrollSaveRequest,
    SalaryHeadCreate
)

from app import crud
from app.excel import generate_excel

router = APIRouter()

# -----------------------------
# High-Speed In-Memory Cache
# -----------------------------
CACHE = {
    "employees": None,
    "employees_ts": 0,
    "salary_heads": None,
    "salary_heads_ts": 0,
    "payroll": {},
}
CACHE_TTL = 300  # 5 minutes

def invalidate_employee_cache():
    CACHE["employees"] = None
    CACHE["employees_ts"] = 0

def invalidate_salary_head_cache():
    CACHE["salary_heads"] = None
    CACHE["salary_heads_ts"] = 0

def invalidate_payroll_cache(month=None, year=None):
    if month is not None and year is not None:
        key = f"{str(month).strip().lower()}_{year}"
        CACHE["payroll"].pop(key, None)
    else:
        CACHE["payroll"].clear()


# -----------------------------
# Employees
# -----------------------------
@router.get(
    "/employees",
    response_model=list[EmployeeResponse],
)
def get_employees(
    db: Session = Depends(get_db),
):
    now = time.time()
    if CACHE["employees"] is not None and (now - CACHE["employees_ts"]) < CACHE_TTL:
        return CACHE["employees"]

    emps = crud.get_employees(db)
    CACHE["employees"] = emps
    CACHE["employees_ts"] = now
    return emps


@router.post("/employees/sanitize-codes")
def sanitize_employee_codes_endpoint(
    db: Session = Depends(get_db),
):
    """
    Scans and auto-fixes any Samarth portal UIDs (VATLHY00090648, etc.) to standard institute codes.
    """
    res = crud.sanitize_all_employee_codes(db)
    invalidate_employee_cache()
    return res


@router.put("/employees/{emp_id}")
def update_employee_endpoint(
    emp_id: int,
    data: Dict[str, Any] = Body(...),
    db: Session = Depends(get_db),
):
    """
    Updates an employee's code, designation, or profile details.
    """
    emp = crud.update_employee(db, emp_id, data)
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")
    invalidate_employee_cache()
    return emp


@router.post("/employees/upload")
async def upload_employees_endpoint(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    """
    Upload and import Employee Master Excel/CSV sheet with 40+ columns.
    """
    try:
        content = await file.read()
        res = crud.upload_employees_from_excel(db, content, file.filename or "employees.xlsx")
        invalidate_employee_cache()
        return res
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to process Employee Master file: {str(e)}")


@router.get("/employees/template")
def download_employee_template():
    """
    Download official blank template with 41 employee columns.
    """
    output = crud.generate_employee_template_excel()
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Employee_Master_Template.xlsx"}
    )


@router.get("/employees/export")
def export_all_employees(
    db: Session = Depends(get_db),
):
    """
    Exports all employees in the database to Excel with complete fields.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    emps = crud.get_employees(db)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employee_Master"

    headers = [
        "employee_id", "title", "first_name", "middle_name", "last_name", "name_in_hindi",
        "employee_type", "nature_of_employment", "organization_unit", "designation",
        "sanctioned_ou", "sanctioned_designation", "guardian_name", "mother_name", "gender",
        "appointed_category", "social_category", "is_pwd", "blood_group", "date_of_birth",
        "date_of_joining", "date_of_superannuation", "mobile_number", "office_phone_number",
        "alternate_mobile_no", "official_email", "personal_email", "residential_address",
        "residential_state", "residential_city", "residential_pincode", "residential_phone_number",
        "permanent_address", "permanent_state", "permanent_city", "permanent_pincode",
        "hometown", "hometown_state", "hometown_city", "hometown_pincode", "status"
    ]
    ws.append(headers)

    for e in emps:
        ws.append([
            e.employee_code, e.title, e.first_name, e.middle_name, e.last_name, e.name_in_hindi,
            e.employee_type, e.nature_of_employment, e.organization_unit, e.designation,
            e.sanctioned_ou, e.sanctioned_designation, e.guardian_name, e.mother_name, e.gender,
            e.appointed_category, e.social_category, e.is_pwd, e.blood_group, e.date_of_birth,
            e.date_of_joining, e.date_of_superannuation, e.mobile_number, e.office_phone_number,
            e.alternate_mobile_no, e.official_email, e.personal_email, e.residential_address,
            e.residential_state, e.residential_city, e.residential_pincode, e.residential_phone_number,
            e.permanent_address, e.permanent_state, e.permanent_city, e.permanent_pincode,
            e.hometown, e.hometown_state, e.hometown_city, e.hometown_pincode, e.status or "Active"
        ])

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align = Alignment(horizontal="center", vertical="center")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=All_Employees_Master.xlsx"}
    )


# -----------------------------
# Salary Heads
# -----------------------------
@router.get(
    "/salary-heads",
    response_model=list[SalaryHeadResponse],
)
def get_salary_heads(
    db: Session = Depends(get_db),
):
    now = time.time()
    if CACHE["salary_heads"] is not None and (now - CACHE["salary_heads_ts"]) < CACHE_TTL:
        return CACHE["salary_heads"]

    heads = crud.get_salary_heads(db)
    CACHE["salary_heads"] = heads
    CACHE["salary_heads_ts"] = now
    return heads

@router.post(
    "/add-salary-heads",
    response_model=SalaryHeadResponse
)
def add_salary_head(
    data: SalaryHeadCreate,
    db: Session = Depends(get_db)
):
    salary_head = data.salary_head.strip()

    if not salary_head:
        raise HTTPException(
            status_code=400,
            detail="Salary head cannot be empty"
        )

    # Check duplicate
    existing = (
        db.query(SalaryHead)
        .filter(SalaryHead.salary_head == salary_head)
        .first()
    )

    if existing:
        raise HTTPException(
            status_code=400,
            detail="Salary head already exists"
        )

    new_salary_head = SalaryHead(
        salary_head=salary_head
    )

    db.add(new_salary_head)
    db.commit()
    db.refresh(new_salary_head)
    invalidate_salary_head_cache()

    return new_salary_head


# -----------------------------
# Payroll Entries
# -----------------------------
@router.get("/payroll")
def get_payroll(
    month: str,
    year: int,
    db: Session = Depends(get_db),
):
    now = time.time()
    key = f"{str(month).strip().lower()}_{year}"
    cached = CACHE["payroll"].get(key)
    if cached and (now - cached[1]) < CACHE_TTL:
        return cached[0]

    rows = crud.get_payroll_entries(db, month, year)
    CACHE["payroll"][key] = (rows, now)
    return rows


# -----------------------------
# Save Payroll
# -----------------------------
@router.post("/payroll/save")
def save_payroll(
    request: PayrollSaveRequest,
    db: Session = Depends(get_db),
):
    crud.save_payroll_entries(
        db,
        request.month,
        request.year,
        request.entries,
    )
    invalidate_payroll_cache(request.month, request.year)

    return {
        "success": True,
        "message": "Payroll saved successfully."
    }


@router.post("/payroll/sync-from-comparison")
def sync_payroll_from_comparison_endpoint(
    data: Dict[str, Any] = Body(...),
    db: Session = Depends(get_db),
):
    """
    1-Click sync of one-time adjustment entries from Salary Comparison into One Time Payroll Entry database.
    """
    month = data.get("month", 6)
    year = data.get("year", 2026)
    entries = data.get("entries", [])
    
    count = crud.sync_payroll_from_comparison(db, int(month), int(year), entries)
    invalidate_payroll_cache(month, year)
    return {
        "success": True,
        "count": count,
        "message": f"Successfully synced {count} one-time adjustment entries into Payroll for month {month}/{year}."
    }


@router.delete("/payroll/delete-all")
def delete_all_payroll_endpoint(
    month: int,
    year: int,
    db: Session = Depends(get_db),
):
    """
    Deletes all one-time payroll entries for the given month and year.
    """
    crud.save_payroll_entries(db, month, year, [])
    invalidate_payroll_cache(month, year)
    return {
        "success": True,
        "message": f"All payroll entries for {month}/{year} deleted successfully."
    }


# -----------------------------
# Export Excel
# -----------------------------

def build_payroll_workbook(rows: list, month: Any, year: Any):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Payroll_Entries"

    # Header styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    headers = [
        "Employee Code",
        "Employee Name",
        "Salary Head",
        "Month",
        "Year",
        "Amount (INR)",
        "Remarks"
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 26

    total_amount = 0.0

    for row_idx, r in enumerate(rows, start=2):
        emp_code = r.get("employee_code") or ""
        emp_name = r.get("employee_name") or ""
        sal_head = r.get("salary_head") or ""
        r_month = r.get("month", month)
        r_year = r.get("year", year)
        
        raw_val = r.get("actual_value") or r.get("amount") or 0
        try:
            val = float(raw_val)
        except Exception:
            val = 0.0
        total_amount += val
        
        remarks = r.get("remarks") or ""

        row_data = [emp_code, emp_name, sal_head, r_month, r_year, val, remarks]
        ws.append(row_data)

        # Apply cell formats
        ws.cell(row=row_idx, column=1).alignment = center_align
        ws.cell(row=row_idx, column=2).alignment = left_align
        ws.cell(row=row_idx, column=3).alignment = left_align
        ws.cell(row=row_idx, column=4).alignment = center_align
        ws.cell(row=row_idx, column=5).alignment = center_align
        
        amount_cell = ws.cell(row=row_idx, column=6)
        amount_cell.alignment = right_align
        amount_cell.number_format = "#,##0.00"
        
        ws.cell(row=row_idx, column=7).alignment = left_align

        for c_i in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=c_i)
            cell.font = data_font
            cell.border = thin_border
            if row_idx % 2 == 1:
                cell.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        ws.row_dimensions[row_idx].height = 20

    # Total Summary Row
    if rows:
        total_row_idx = len(rows) + 2
        ws.cell(row=total_row_idx, column=1, value="")
        ws.cell(row=total_row_idx, column=2, value="TOTAL").font = bold_font
        ws.cell(row=total_row_idx, column=3, value=f"{len(rows)} Entries").font = bold_font
        ws.cell(row=total_row_idx, column=4, value="")
        ws.cell(row=total_row_idx, column=5, value="")
        
        tot_cell = ws.cell(row=total_row_idx, column=6, value=total_amount)
        tot_cell.font = bold_font
        tot_cell.alignment = right_align
        tot_cell.number_format = "#,##0.00"
        
        ws.cell(row=total_row_idx, column=7, value="")

        total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        for c_i in range(1, len(headers) + 1):
            c_elem = ws.cell(row=total_row_idx, column=c_i)
            c_elem.fill = total_fill
            c_elem.border = thin_border
        ws.row_dimensions[total_row_idx].height = 22

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    excel = BytesIO()
    wb.save(excel)
    excel.seek(0)
    return excel


@router.get("/payroll/export")
def export_excel(
    month: str,
    year: int,
    db: Session = Depends(get_db),
):
    rows = crud.get_payroll_entries(db, month, year)
    excel = build_payroll_workbook(rows, month, year)
    filename = f"Payroll_{month}_{year}.xlsx"

    return StreamingResponse(
        excel,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        },
    )


@router.post("/payroll/export-rows")
def export_payroll_rows_endpoint(
    payload: Dict[str, Any] = Body(...),
):
    """
    Directly exports given payroll rows into styled Excel spreadsheet.
    """
    month = payload.get("month", "01")
    year = payload.get("year", 2026)
    rows = payload.get("rows", [])

    excel = build_payroll_workbook(rows, month, year)
    filename = f"Payroll_Entries_{month}_{year}.xlsx"

    return StreamingResponse(
        excel,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        },
    )


# -----------------------------
# Salary Comparison Module
# -----------------------------
import json
from fastapi import UploadFile, File, Form, Body
from app.comparison.excel_parser import parse_excel_file
from app.comparison.comparator import compare_monthly_datasets
from app.comparison.exporter import generate_comparison_excel_report
from app.comparison.sample_data import get_sample_comparison_dataset, generate_sample_month_workbook


@router.post("/salary-comparison/compare")
async def compare_salary_files(
    files: list[UploadFile] = File(...),
    month_labels: Optional[str] = Form(None),
):
    """
    Accepts 2, 3, or more uploaded Excel files (.xlsx, .xls).
    month_labels can be passed as JSON string or comma-separated string.
    Returns complete structured comparison report.
    """
    if len(files) < 2:
        raise HTTPException(
            status_code=400,
            detail="Please upload at least 2 monthly salary Excel files for comparison."
        )

    labels = []
    if month_labels:
        try:
            parsed_labels = json.loads(month_labels)
            if isinstance(parsed_labels, list):
                labels = [str(l).strip() for l in parsed_labels]
        except Exception:
            labels = [l.strip() for l in month_labels.split(",") if l.strip()]

    # Fallback to filenames or Month 1, Month 2...
    while len(labels) < len(files):
        idx = len(labels)
        file_name = files[idx].filename or f"Month {idx + 1}"
        # Clean up file extension from label if using filename
        clean_name = file_name.rsplit(".", 1)[0] if "." in file_name else file_name
        labels.append(clean_name or f"Month {idx + 1}")

    monthly_datasets = []

    for idx, uploaded_file in enumerate(files):
        content = await uploaded_file.read()
        if not content:
            raise HTTPException(
                status_code=400,
                detail=f"File '{uploaded_file.filename}' is empty."
            )

        try:
            category_data = parse_excel_file(content, uploaded_file.filename or f"File_{idx+1}")
            if not category_data:
                raise HTTPException(
                    status_code=400,
                    detail=f"No valid salary sheets or columns found in '{uploaded_file.filename}'."
                )

            monthly_datasets.append({
                "month_index": idx,
                "month_label": labels[idx],
                "categories": category_data
            })
        except HTTPException:
            raise
        except Exception as err:
            raise HTTPException(
                status_code=400,
                detail=f"Error parsing Excel file '{uploaded_file.filename}': {str(err)}"
            )

    try:
        comparison_result = compare_monthly_datasets(monthly_datasets)
        return comparison_result
    except Exception as err:
        raise HTTPException(
            status_code=500,
            detail=f"Comparison calculation error: {str(err)}"
        )


@router.post("/salary-comparison/export")
def export_salary_comparison_excel(
    payload: dict = Body(...),
):
    """
    Accepts comparison report data and returns downloadable multi-sheet .xlsx file.
    """
    try:
        excel_stream = generate_comparison_excel_report(payload)
        months = payload.get("months", ["Report"])
        filename = f"Salary_Comparison_{'_vs_'.join([m.replace(' ', '_') for m in months])}.xlsx"

        return StreamingResponse(
            excel_stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"'
            },
        )
    except Exception as err:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to generate comparison Excel report: {str(err)}"
        )


@router.get("/salary-comparison/sample")
def get_sample_comparison(
    month: Optional[int] = 6,
    year: Optional[int] = 2026,
):
    """
    Returns sample 3-month comparison dataset for immediate UI preview and demo based on selected month/year.
    """
    try:
        return get_sample_comparison_dataset(current_year=year or 2026, current_month=month or 6)
    except Exception as err:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to generate sample comparison data: {str(err)}"
        )


@router.get("/salary-comparison/sample-file/{month_index}")
def download_sample_excel_file(
    month_index: int,
    month: Optional[int] = 6,
    year: Optional[int] = 2026,
):
    """
    Download individual sample monthly Excel file for the requested month slot (0=Month N-2, 1=Month N-1, 2=Month N).
    """
    from app.comparison.comparator import get_comparison_months
    months_meta = get_comparison_months(year or 2026, month or 6, count=3)
    if month_index < 0 or month_index >= len(months_meta):
        raise HTTPException(status_code=400, detail=f"Invalid month index (must be 0, 1, or 2).")

    file_bytes = generate_sample_month_workbook(month_index)
    month_label = months_meta[month_index]["label"].replace(" ", "_")
    filename = f"Salary_{month_label}.xlsx"

    return StreamingResponse(
        io.BytesIO(file_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        },
    )


@router.post("/salary-comparison/reconcile-samarth")
async def reconcile_samarth_endpoint(
    samarth_file: UploadFile = File(...),
    current_month_file: Optional[UploadFile] = File(None),
    month_label: Optional[str] = Form("Current Month"),
):
    """
    Reconciles an uploaded Samarth Generated Excel sheet directly against the Current Month Salary Sheet.
    """
    from app.comparison.excel_parser import parse_excel_file
    from app.comparison.comparator import reconcile_current_with_samarth
    from app.comparison.sample_data import generate_sample_month_workbook

    try:
        samarth_bytes = await samarth_file.read()
        samarth_parsed = parse_excel_file(samarth_bytes, samarth_file.filename or "samarth.xlsx")

        if current_month_file is not None:
            current_bytes = await current_month_file.read()
            current_parsed = parse_excel_file(current_bytes, current_month_file.filename or "current_salary.xlsx")
        else:
            # Fallback to generated current month sample dataset
            current_bytes = generate_sample_month_workbook(1)
            current_parsed = parse_excel_file(current_bytes, "current_salary.xlsx")

        return reconcile_current_with_samarth(
            current_categories=current_parsed,
            samarth_categories=samarth_parsed,
            month_label=month_label or "Current Month"
        )
    except Exception as err:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to reconcile Samarth Generated sheet: {str(err)}"
        )


@router.post("/salary-comparison/compare-current-vs-samarth")
async def compare_current_vs_samarth_endpoint(
    current_file: UploadFile = File(...),
    samarth_file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    """
    Compares TWO Excel files directly:
    1) Current Month Salary Excel
    2) Samarth Generated Salary Excel
    """
    from app.comparison.excel_parser import parse_excel_file
    from app.comparison.current_vs_samarth import compare_current_vs_samarth

    try:
        curr_bytes = await current_file.read()
        sam_bytes = await samarth_file.read()

        curr_parsed = parse_excel_file(curr_bytes, current_file.filename or "Current_Salary.xlsx")
        sam_parsed = parse_excel_file(sam_bytes, samarth_file.filename or "Samarth_Salary.xlsx")

        if not curr_parsed or not any(len(c.get("employees", [])) > 0 for c in curr_parsed.values()):
            raise HTTPException(status_code=400, detail="Current Month Excel file contains no readable employee records.")

        if not sam_parsed or not any(len(c.get("employees", [])) > 0 for c in sam_parsed.values()):
            raise HTTPException(status_code=400, detail="Samarth Generated Excel file contains no readable employee records.")

        return compare_current_vs_samarth(
            current_categories=curr_parsed,
            samarth_categories=sam_parsed,
            current_filename=current_file.filename or "Current Month Salary Report",
            samarth_filename=samarth_file.filename or "Samarth Generated Salary Report",
            db=db
        )
    except HTTPException:
        raise
    except Exception as err:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to process Current vs Samarth salary comparison: {str(err)}"
        )


@router.get("/salary-comparison/sample-current-vs-samarth")
def sample_current_vs_samarth_endpoint():
    """
    Generates and returns immediate sample comparison result between Current Month Sheet & Samarth Sheet.
    """
    from app.comparison.sample_data import generate_sample_month_workbook
    from app.comparison.excel_parser import parse_excel_file
    from app.comparison.current_vs_samarth import compare_current_vs_samarth

    try:
        curr_bytes = generate_sample_month_workbook(1)
        sam_bytes = generate_sample_month_workbook(2)

        curr_parsed = parse_excel_file(curr_bytes, "June_2026_Salary_Report.xlsx")
        sam_parsed = parse_excel_file(sam_bytes, "June_2026_Samarth_Generated.xlsx")

        return compare_current_vs_samarth(
            current_categories=curr_parsed,
            samarth_categories=sam_parsed,
            current_filename="June 2026 Salary Report.xlsx",
            samarth_filename="June 2026 Samarth Salary Generated.xlsx"
        )
    except Exception as err:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to generate sample Current vs Samarth comparison: {str(err)}"
        )


@router.post("/salary-comparison/export-current-vs-samarth")
async def export_current_vs_samarth_endpoint(
    comparison_data: Dict[str, Any] = Body(...)
):
    """
    Generates and downloads styled Excel report for Current Month vs Samarth Comparison.
    """
    from app.comparison.current_vs_samarth import generate_current_vs_samarth_excel

    try:
        excel_stream = generate_current_vs_samarth_excel(comparison_data)
        filename = "Current_Month_vs_Samarth_Salary_Comparison.xlsx"

        return StreamingResponse(
            excel_stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"'
            },
        )
    except Exception as err:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to generate Current vs Samarth Excel export: {str(err)}"
        )


@router.post("/salary-comparison/export-one-time-entries")
async def export_one_time_entries_endpoint(
    comparison_data: Dict[str, Any] = Body(...)
):
    """
    Generates and downloads dedicated One-Time Payroll Entries Excel workbook.
    Excludes Gross / Total Earnings, Total Deductions, and Net Pay totals.
    """
    from app.comparison.current_vs_samarth import generate_one_time_entries_excel

    try:
        excel_stream = generate_one_time_entries_excel(comparison_data)
        filename = "One_Time_Payroll_Entries_Adjustments.xlsx"

        return StreamingResponse(
            excel_stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"'
            },
        )
    except Exception as err:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to generate One-Time Entries Excel export: {str(err)}"
        )