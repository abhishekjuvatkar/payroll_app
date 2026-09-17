import io
import re
from typing import Dict, List, Any, Optional, Tuple, Set
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.comparison.excel_parser import parse_excel_file, normalize_text, normalize_employee_name, CATEGORY_NAMES

# Exporter styles
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Arial", size=14, bold=True, color="1E293B")
SUBTITLE_FONT = Font(name="Arial", size=10, italic=True, color="64748B")
BOLD_FONT = Font(name="Arial", size=10, bold=True, color="0F172A")
NORMAL_FONT = Font(name="Arial", size=10, color="0F172A")

HEADER_FILL_PRIMARY = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid") # Deep Blue
HEADER_FILL_SECONDARY = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid") # Teal
HEADER_FILL_ACCENT = PatternFill(start_color="475569", end_color="475569", fill_type="solid") # Slate

FILL_MATCH = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Soft green
FILL_MISMATCH = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Soft red
FILL_NOT_AVAIL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Soft yellow
FILL_TOTAL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid") # Slate light

THIN_BORDER = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

CURRENCY_FORMAT = '0;[Red](-0);"-"'
INTEGER_FORMAT = '0'

def fmt_clean_num(v: Any) -> str:
    if v is None:
        return "0"
    try:
        f = float(v)
        if round(f, 2) == round(f):
            return str(int(round(f)))
        return f"{round(f, 2):.2f}".rstrip("0").rstrip(".")
    except Exception:
        return str(v)

def to_clean_num(v: float) -> Any:
    if round(v, 2) == round(v):
        return int(round(v))
    return round(v, 2)

AUTO_CALCULATED_HEAD_KEYS = {
    "BASIC", "DA", "HRA", "TA", "DA_ON_TA",
    "ARREARS_SALARY", "ARREARS",
    "TOTAL_EARNINGS", "TOTAL_DEDUCTIONS", "NET_PAY",
    "GROSS_WITH_EMPLOYER", "DEDUCTION_WITH_EMPLOYER"
}

AUTO_CALCULATED_NAMES = {
    "BASIC PAY", "BASIC",
    "DEARNESS ALLOWANCE", "DEARNESS ALLOWANCE (DA)", "DA",
    "HOUSE RENT ALLOWANCE", "HOUSE RENT ALLOWANCE (HRA)", "HRA",
    "TRANSPORT ALLOWANCE", "TRANSPORT ALLOWANCE (TA)", "TA", "TPTA",
    "DA ON TA", "DA ON TPTA", "DA(TA)", "DA ON TRANSPORT ALLOWANCE",
    "ARREARS ON SALARY", "ARREARS", "ARREAR ON SALARY", "ARREAR",
    "GROSS / TOTAL EARNINGS", "TOTAL EARNINGS", "GROSS", "GROSS SALARY",
    "TOTAL DEDUCTIONS", "DEDUCTIONS", "NET PAY", "NET AMOUNT",
    "GROSS WITH EMPLOYER", "DEDUCTION WITH EMPLOYER"
}

def is_auto_calculated_head(canonical_k: str, head_n: str) -> bool:
    if (canonical_k or "").upper() in AUTO_CALCULATED_HEAD_KEYS:
        return True
    u_n = (head_n or "").strip().upper()
    if u_n in AUTO_CALCULATED_NAMES:
        return True
    clean = re.sub(r"\s*\([^)]*\)", "", u_n).strip()
    if clean in AUTO_CALCULATED_NAMES:
        return True
    return False

import itertools

def get_name_variations(name: str) -> Set[str]:
    """
    Generates unambiguous permutations, initials, and normalized aliases for Indian names.
    Guarantees that pure initials (like 'A G', 'G A') are NEVER generated, preventing cross-employee collisions.
    e.g. 'Ravikumar Chettiannan' -> 'RAVIKUMAR CHETTIANNAN', 'CHETTIANNAN RAVIKUMAR', 'C RAVIKUMAR', 'RAVIKUMAR C'
    e.g. 'Hemanth Kumar Chinthapalli' -> 'HEMANTH KUMAR CHINTHAPALLI', 'CHINTHAPALLI HEMANTH KUMAR', 'C HEMANTH KUMAR', etc.
    """
    if not name:
        return set()

    clean = re.sub(r"^(dr\.|prof\.|mr\.|mrs\.|ms\.|shri|smt\.)\s+", "", name.strip(), flags=re.IGNORECASE)
    clean = re.sub(r"[^A-Za-z0-9\s]+", " ", clean)
    clean = re.sub(r"\s+", " ", clean).strip().upper()
    
    if not clean:
        return set()

    tokens = [t for t in clean.split() if t]
    if not tokens:
        return set()

    variations = set()
    variations.add(clean)

    if len(tokens) == 1:
        return variations

    # Sorted tokens
    sorted_tokens = " ".join(sorted(tokens))
    variations.add(sorted_tokens)

    # Permutations for 2 or 3 tokens (e.g. Ravikumar Chettiannan -> Chettiannan Ravikumar)
    if len(tokens) <= 3:
        for p in itertools.permutations(tokens):
            variations.add(" ".join(p))

    # Single-letter token removal (e.g. 'AMAR K GAONKAR' -> 'AMAR GAONKAR')
    multi_char_tokens = [t for t in tokens if len(t) > 1]
    if len(multi_char_tokens) >= 2:
        variations.add(" ".join(multi_char_tokens))
        variations.add(" ".join(sorted(multi_char_tokens)))
        if len(multi_char_tokens) <= 3:
            for p in itertools.permutations(multi_char_tokens):
                variations.add(" ".join(p))

    # Initial variations: ONLY 1 initial combined with full word tokens (len >= 3)
    # NEVER generate pure initials (e.g. 'A G' or 'G A')
    if len(tokens) >= 2:
        rest_1 = " ".join(tokens[1:])
        rest_0 = " ".join(tokens[:-1])
        first_init = f"{tokens[0][0]} {rest_1}"
        last_init = f"{rest_0} {tokens[-1][0]}"
        init_first_inv = f"{rest_1} {tokens[0][0]}"
        init_last_inv = f"{tokens[-1][0]} {rest_0}"
        
        for cand in [first_init, last_init, init_first_inv, init_last_inv]:
            parts = cand.split()
            if any(len(p) >= 3 for p in parts) and len(cand) >= 5:
                variations.add(cand)

    # Strict filter: Must have at least 2 parts, total length >= 5, and at least one word >= 3 chars
    return {v for v in variations if len(v) >= 5 and len(v.split()) >= 2 and any(len(p) >= 3 for p in v.split())}

def clean_employee_lookup_key(name: str) -> str:
    """
    Strips titles, dots, single-letter middle initials and extra spaces.
    e.g. 'Dr. Dhiraj V. Patil' -> 'DHIRAJ PATIL'
         'Dhiraj Patil' -> 'DHIRAJ PATIL'
         'DHIRAJ V PATIL' -> 'DHIRAJ PATIL'
    """
    if not name:
        return ""
    clean = normalize_employee_name(name)
    # Remove single letter middle initials with or without dots (e.g., ' V. ', ' V ')
    clean = re.sub(r"\b[A-Z]\b\.?", "", clean)
    clean = re.sub(r"[^A-Z0-9]+", " ", clean).strip()
    return re.sub(r"\s+", " ", clean)

def compare_current_vs_samarth(
    current_categories: Dict[str, Any],
    samarth_categories: Dict[str, Any],
    current_filename: str = "Current Month Salary Report",
    samarth_filename: str = "Samarth Generated Salary Report",
    db: Optional[Any] = None
) -> Dict[str, Any]:
    """
    Compares ONLY TWO files:
    1) Current Month Salary Excel
    2) Samarth Generated Salary Excel
    
    Formula: Difference = Current Month Value - Samarth Value
    Alphabetical employee order A -> Z.
    Excludes Totals from Mismatch counts and One-Time Adjustment entries.
    Matches with official Employee Master database records.
    Ensures ALL heads from Current Month are compared and never ignored.
    """
    all_cat_keys = set(current_categories.keys()).union(set(samarth_categories.keys()))
    
    # 0. Lookup official Employee Master codes from database with strict collision prevention
    db_emp_codes: Dict[str, str] = {}
    if db is not None:
        try:
            from app.models import Employee
            collisions = set()
            all_db_emps = db.query(Employee).all()
            for e in all_db_emps:
                code = (e.employee_code or "").strip().upper()
                raw_name = (e.employee_name or "").strip().upper()
                norm_e = normalize_employee_name(e.employee_name)
                clean_e = clean_employee_lookup_key(e.employee_name)

                if code and code != raw_name:
                    db_emp_codes[code] = code
                    if raw_name:
                        db_emp_codes[raw_name] = code
                    if norm_e:
                        db_emp_codes[norm_e] = code
                    if clean_e and len(clean_e.split()) >= 2:
                        db_emp_codes[clean_e] = code

                    for var in get_name_variations(e.employee_name):
                        if var:
                            if var in db_emp_codes and db_emp_codes[var] != code:
                                collisions.add(var)
                            else:
                                db_emp_codes[var] = code
                                
            # Discard any ambiguous variations that matched multiple employees
            for c in collisions:
                db_emp_codes.pop(c, None)
        except Exception:
            pass

    # 1. Build unified employee registry & all known columns
    current_emp_registry: Dict[str, Dict[str, Any]] = {}
    samarth_emp_registry: Dict[str, Dict[str, Any]] = {}
    all_unique_employees: Dict[str, Dict[str, Any]] = {}
    all_known_columns: Dict[str, Dict[str, Any]] = {}

    def resolve_emp_code(raw_id: str, raw_name: str, norm_n: str, clean_n: str) -> str:
        clean_raw_id = (raw_id or "").strip().upper()
        if clean_raw_id:
            if clean_raw_id in db_emp_codes:
                return db_emp_codes[clean_raw_id]
            if re.match(r"^[A-Z]{1,4}\d{2,6}[A-Z]?$", clean_raw_id):
                return clean_raw_id

        for var in get_name_variations(raw_name):
            if var in db_emp_codes:
                return db_emp_codes[var]

        for var in get_name_variations(norm_n):
            if var in db_emp_codes:
                return db_emp_codes[var]

        return (
            db_emp_codes.get(clean_n) or
            db_emp_codes.get(raw_name.strip().upper()) or
            (clean_raw_id if clean_raw_id and clean_raw_id != raw_name.strip().upper() else clean_raw_id)
        )
    
    for cat_key in all_cat_keys:
        curr_cat = current_categories.get(cat_key, {})
        for col in curr_cat.get("columns", []):
            all_known_columns[col["key"]] = col

        for emp in curr_cat.get("employees", []):
            emp_id = emp.get("employee_id") or ""
            raw_name = emp.get("employee_name") or ""
            norm_name = emp.get("normalized_name") or normalize_employee_name(raw_name)
            clean_name = clean_employee_lookup_key(raw_name)
            
            official_id = resolve_emp_code(emp_id, raw_name, norm_name, clean_name) or emp_id
            emp_info = {
                **emp,
                "employee_id": official_id,
                "normalized_name": norm_name,
                "clean_name": clean_name,
                "category_key": cat_key,
                "category_name": curr_cat.get("category_name", "Staff")
            }
            
            if official_id:
                current_emp_registry[official_id] = emp_info
            current_emp_registry[norm_name] = emp_info
            if clean_name:
                current_emp_registry[clean_name] = emp_info
            current_emp_registry[raw_name.strip().upper()] = emp_info
            for var in get_name_variations(raw_name):
                current_emp_registry[var] = emp_info
            
            # Use official_id or clean_name for deduplication key
            primary_key = official_id if (official_id and official_id != raw_name.strip().upper()) else (clean_name or norm_name)
            if primary_key not in all_unique_employees:
                all_unique_employees[primary_key] = emp_info

        sam_cat = samarth_categories.get(cat_key, {})
        for col in sam_cat.get("columns", []):
            all_known_columns[col["key"]] = col

        for emp in sam_cat.get("employees", []):
            emp_id = emp.get("employee_id") or ""
            raw_name = emp.get("employee_name") or ""
            norm_name = emp.get("normalized_name") or normalize_employee_name(raw_name)
            clean_name = clean_employee_lookup_key(raw_name)
            
            official_id = resolve_emp_code(emp_id, raw_name, norm_name, clean_name) or emp_id
            emp_info = {
                **emp,
                "employee_id": official_id,
                "normalized_name": norm_name,
                "clean_name": clean_name,
                "category_key": cat_key,
                "category_name": sam_cat.get("category_name", "Staff")
            }
            
            if official_id:
                samarth_emp_registry[official_id] = emp_info
            samarth_emp_registry[norm_name] = emp_info
            if clean_name:
                samarth_emp_registry[clean_name] = emp_info
            samarth_emp_registry[raw_name.strip().upper()] = emp_info
            for var in get_name_variations(raw_name):
                samarth_emp_registry[var] = emp_info
            
            primary_key = official_id if (official_id and official_id != raw_name.strip().upper()) else (clean_name or norm_name)
            if primary_key not in all_unique_employees:
                all_unique_employees[primary_key] = emp_info
            else:
                # Merge Samarth metadata
                existing = all_unique_employees[primary_key]
                if not db_emp_codes.get(norm_name) and official_id and official_id != norm_name:
                    existing["employee_id"] = official_id
                if emp.get("employee_name") and not emp.get("employee_name").isupper():
                    existing["employee_name"] = emp["employee_name"]

    # 2. Collect all distinct columns (salary heads) across categories
    category_columns_map: Dict[str, Dict[str, Dict[str, Any]]] = {}
    for cat_key in all_cat_keys:
        category_columns_map[cat_key] = {}
        curr_cols = current_categories.get(cat_key, {}).get("columns", [])
        sam_cols = samarth_categories.get(cat_key, {}).get("columns", [])
        for c in curr_cols:
            category_columns_map[cat_key][c["key"]] = c
            all_known_columns[c["key"]] = c
        for c in sam_cols:
            category_columns_map[cat_key][c["key"]] = c
            all_known_columns[c["key"]] = c

    # 3. Sort employees alphabetically (A -> Z) case-insensitive
    unique_emp_list = list(all_unique_employees.values())
    unique_emp_list.sort(key=lambda e: (e["employee_name"].strip().upper(), e.get("employee_id", "")))

    employee_grouped_list = []
    flat_rows = []
    one_time_entries = []
    
    total_matches = 0
    total_mismatches = 0
    total_missing_values = 0
    matched_employees_count = 0
    missing_in_current_count = 0
    missing_in_samarth_count = 0
    
    total_current_amount = 0.0
    total_samarth_amount = 0.0
    
    head_mismatch_stats: Dict[str, Dict[str, Any]] = {}
    
    # 4. Compare employee by employee
    for emp_data in unique_emp_list:
        emp_name = emp_data["employee_name"]
        emp_id = emp_data["employee_id"]
        norm_name = emp_data.get("normalized_name") or normalize_employee_name(emp_name)
        clean_name = emp_data.get("clean_name") or clean_employee_lookup_key(emp_name)
        raw_upper = emp_name.strip().upper()
        cat_key = emp_data.get("category_key", "staff")
        cat_name = emp_data.get("category_name", "Staff")
        
        # Robust lookup across ID, variations, exact normalized name, clean phonetic name, and raw upper name
        curr_emp = current_emp_registry.get(emp_id)
        if not curr_emp:
            for var in get_name_variations(emp_name):
                if var in current_emp_registry:
                    curr_emp = current_emp_registry[var]
                    break
        if not curr_emp:
            curr_emp = (
                current_emp_registry.get(norm_name) or
                current_emp_registry.get(clean_name) or
                current_emp_registry.get(raw_upper)
            )

        sam_emp = samarth_emp_registry.get(emp_id)
        if not sam_emp:
            for var in get_name_variations(emp_name):
                if var in samarth_emp_registry:
                    sam_emp = samarth_emp_registry[var]
                    break
        if not sam_emp:
            sam_emp = (
                samarth_emp_registry.get(norm_name) or
                samarth_emp_registry.get(clean_name) or
                samarth_emp_registry.get(raw_upper)
            )
        
        # Determine employee match status
        if curr_emp and sam_emp:
            emp_match_status = "MATCHED"
            matched_employees_count += 1
        elif sam_emp and not curr_emp:
            emp_match_status = "MISSING_IN_CURRENT"
            missing_in_current_count += 1
        else:
            emp_match_status = "MISSING_IN_SAMARTH"
            missing_in_samarth_count += 1
            
        emp_heads_list = []
        emp_current_total = 0.0
        emp_samarth_total = 0.0
        emp_mismatch_count = 0
        
        # Combine category columns with all heads present in this specific employee's Current Month & Samarth data
        cat_cols = category_columns_map.get(cat_key, {})
        all_emp_head_keys = list(cat_cols.keys())
        
        if curr_emp:
            for k in curr_emp.get("values", {}).keys():
                if k not in all_emp_head_keys:
                    all_emp_head_keys.append(k)
                    
        if sam_emp:
            for k in sam_emp.get("values", {}).keys():
                if k not in all_emp_head_keys:
                    all_emp_head_keys.append(k)

        from app.comparison.excel_parser import map_salary_head_to_canonical

        for col_key in all_emp_head_keys:
            col_meta = cat_cols.get(col_key) or all_known_columns.get(col_key)
            if not col_meta:
                _, d_name, is_tot, is_emp = map_salary_head_to_canonical(col_key)
                col_meta = {
                    "key": col_key,
                    "name": d_name,
                    "is_total": is_tot,
                    "is_employer_contribution": is_emp
                }

            head_name = col_meta["name"]
            is_total = col_meta.get("is_total", False)
            is_emp_contrib = col_meta.get("is_employer_contribution", False)
            
            has_in_current = curr_emp is not None and col_key in curr_emp["values"]
            has_in_samarth = sam_emp is not None and col_key in sam_emp["values"]
            
            curr_val = curr_emp["values"].get(col_key, 0.0) if curr_emp else 0.0
            sam_val = sam_emp["values"].get(col_key, 0.0) if sam_emp else 0.0
            
            # Difference formula: Current Month Value - Samarth Value
            diff = round(curr_val - sam_val, 2)
            
            if not has_in_current and has_in_samarth:
                if abs(sam_val) >= 0.01:
                    status = "NOT_AVAILABLE_IN_CURRENT"
                    total_missing_values += 1
                else:
                    status = "MATCH"
                    if not is_total:
                        total_matches += 1
            elif has_in_current and not has_in_samarth:
                if abs(curr_val) >= 0.01:
                    status = "NOT_AVAILABLE_IN_SAMARTH"
                    total_missing_values += 1
                else:
                    status = "MATCH"
                    if not is_total:
                        total_matches += 1
            elif abs(diff) < 0.01:
                status = "MATCH"
                if not is_total:
                    total_matches += 1
            else:
                status = "MISMATCH"

            # Count and track discrepancy stats
            has_discrepancy = (status in ["MISMATCH", "NOT_AVAILABLE_IN_SAMARTH", "NOT_AVAILABLE_IN_CURRENT"]) and (abs(diff) >= 0.01)
            
            if has_discrepancy and not is_total:
                emp_mismatch_count += 1
                total_mismatches += 1
                
                # Track in head mismatch summary (EXCLUDING totals)
                if head_name not in head_mismatch_stats:
                    head_mismatch_stats[head_name] = {
                        "salary_head": head_name,
                        "category": cat_name,
                        "employees_changed": 0,
                        "total_difference": 0.0
                    }
                head_mismatch_stats[head_name]["employees_changed"] += 1
                head_mismatch_stats[head_name]["total_difference"] = round(
                    head_mismatch_stats[head_name]["total_difference"] + diff, 2
                )

                # Generate One-Time Entry record ONLY if:
                # 1. Current Month has positive non-zero value (curr_val >= 0.01)
                # 2. It is NOT an auto-calculated recurring head (DA, DA on TA, HRA, TA, Basic Pay)
                # 3. It is NOT a calculated aggregate total
                if curr_val >= 0.01 and not is_auto_calculated_head(col_key, head_name) and not is_total:
                    entry_type = "Deduction Adjustment" if any(w in head_name.upper() for w in ["TAX", "DED", "NPS", "FEE", "CHARGES", "GPF", "RECOVERY", "CLUB"]) else "Earning Adjustment"
                    if is_emp_contrib:
                        entry_type = "Employer Cost Adjustment"
                        
                    if status == "NOT_AVAILABLE_IN_SAMARTH":
                        rem = f"One-time adjustment for {head_name}: present in Current Month ({fmt_clean_num(curr_val)}) but missing in Samarth"
                        direction = "Current > Samarth (+)"
                    else:
                        rem = f"One-time adjustment for {head_name}: variance of {fmt_clean_num(abs(diff))}"
                        direction = "Current > Samarth (+)" if diff > 0 else "Samarth > Current (-)"

                    one_time_entries.append({
                        "s_no": len(one_time_entries) + 1,
                        "employee_id": emp_id,
                        "employee_name": emp_name,
                        "category": cat_name,
                        "salary_head": head_name,
                        "entry_type": entry_type,
                        "current_value": to_clean_num(curr_val),
                        "samarth_value": to_clean_num(sam_val),
                        "difference": to_clean_num(diff),
                        "adjustment_amount": to_clean_num(curr_val),
                        "adjustment_direction": direction,
                        "remarks": rem
                    })
                
            if not is_total and not is_emp_contrib:
                emp_current_total += curr_val
                emp_samarth_total += sam_val
                total_current_amount += curr_val
                total_samarth_amount += sam_val
                
            head_item = {
                "salary_head": head_name,
                "head_key": col_key,
                "current_value": to_clean_num(curr_val),
                "samarth_value": to_clean_num(sam_val),
                "difference": to_clean_num(diff),
                "status": status,
                "is_total": is_total,
                "is_employer_contribution": is_emp_contrib
            }
            emp_heads_list.append(head_item)
            
            # Add to flat list for table view
            flat_rows.append({
                "id": f"{emp_id}_{col_key}",
                "employee_name": emp_name,
                "employee_id": emp_id,
                "category": cat_name,
                "salary_head": head_name,
                "head_key": col_key,
                "current_value": to_clean_num(curr_val),
                "samarth_value": to_clean_num(sam_val),
                "difference": to_clean_num(diff),
                "status": status,
                "is_total": is_total,
                "is_employer_contribution": is_emp_contrib
            })
            
        emp_diff = round(emp_current_total - emp_samarth_total, 2)
        employee_grouped_list.append({
            "employee_name": emp_name,
            "employee_id": emp_id,
            "category": cat_name,
            "match_status": emp_match_status,
            "current_total": to_clean_num(emp_current_total),
            "samarth_total": to_clean_num(emp_samarth_total),
            "difference": to_clean_num(emp_diff),
            "mismatch_count": emp_mismatch_count,
            "heads": emp_heads_list
        })

    # Sort head mismatch summary by number of discrepancies descending
    head_mismatch_list = list(head_mismatch_stats.values())
    for hm in head_mismatch_list:
        hm["total_difference"] = to_clean_num(hm["total_difference"])
    head_mismatch_list.sort(key=lambda h: (-h["employees_changed"], h["salary_head"]))

    # Summary KPIs
    total_diff_overall = round(total_current_amount - total_samarth_amount, 2)
    summary_kpi = {
        "total_employees": len(unique_emp_list),
        "matched_employees": matched_employees_count,
        "missing_in_current": missing_in_current_count,
        "missing_in_samarth": missing_in_samarth_count,
        "total_salary_heads_compared": len(category_columns_map.get("staff", {})) or len(head_mismatch_list),
        "total_matches": total_matches,
        "total_mismatches": total_mismatches,
        "total_one_time_entries": len(one_time_entries),
        "total_missing_values": total_missing_values,
        "total_current_amount": to_clean_num(total_current_amount),
        "total_samarth_amount": to_clean_num(total_samarth_amount),
        "total_difference": to_clean_num(total_diff_overall),
        "current_filename": current_filename,
        "samarth_filename": samarth_filename
    }

    return {
        "summary": summary_kpi,
        "head_mismatch_summary": head_mismatch_list,
        "employee_grouped_comparison": employee_grouped_list,
        "one_time_entries": one_time_entries,
        "flat_comparison_rows": flat_rows,
        "categories": [CATEGORY_NAMES.get(k, k) for k in all_cat_keys]
    }

def auto_fit_columns(ws, min_width=12, max_width=45):
    """Auto-fit column widths."""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if '\n' in val_str:
                val_str = max(val_str.split('\n'), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 4, max_width))

def style_header_row(ws, row_idx: int, fill=HEADER_FILL_PRIMARY):
    for cell in ws[row_idx]:
        cell.font = HEADER_FONT
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

def generate_current_vs_samarth_excel(comparison_result: Dict[str, Any]) -> io.BytesIO:
    """
    Generates a dedicated Excel report for Current Month vs Samarth Comparison.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet
    
    summary = comparison_result.get("summary", {})
    grouped_emps = comparison_result.get("employee_grouped_comparison", [])
    head_mismatches = comparison_result.get("head_mismatch_summary", [])
    flat_rows = comparison_result.get("flat_comparison_rows", [])

    # -------------------------------------------------------------
    # 1. SHEET: Summary
    # -------------------------------------------------------------
    ws_sum = wb.create_sheet(title="Summary")
    ws_sum.views.sheetView[0].showGridLines = True
    
    ws_sum["A1"] = "CURRENT MONTH vs SAMARTH SALARY COMPARISON REPORT"
    ws_sum["A1"].font = TITLE_FONT
    ws_sum["A2"] = f"Files Compared: {summary.get('current_filename', 'Current Month')}  VS  {summary.get('samarth_filename', 'Samarth Generated')}"
    ws_sum["A2"].font = SUBTITLE_FONT

    kpi_rows = [
        ("Total Unique Employees", summary.get("total_employees", 0), INTEGER_FORMAT),
        ("Matched in Both Files", summary.get("matched_employees", 0), INTEGER_FORMAT),
        ("Missing in Current Month File", summary.get("missing_in_current", 0), INTEGER_FORMAT),
        ("Missing in Samarth File", summary.get("missing_in_samarth", 0), INTEGER_FORMAT),
        ("Total Salary Head Matches", summary.get("total_matches", 0), INTEGER_FORMAT),
        ("Total Salary Head Mismatches", summary.get("total_mismatches", 0), INTEGER_FORMAT),
        ("Total Missing Head Values", summary.get("total_missing_values", 0), INTEGER_FORMAT),
        ("Total Current Month Amount (₹)", summary.get("total_current_amount", 0.0), CURRENCY_FORMAT),
        ("Total Samarth Amount (₹)", summary.get("total_samarth_amount", 0.0), CURRENCY_FORMAT),
        ("Net Total Difference (₹)", summary.get("total_difference", 0.0), CURRENCY_FORMAT),
    ]

    ws_sum.cell(row=4, column=1, value="Metric").font = BOLD_FONT
    ws_sum.cell(row=4, column=2, value="Value").font = BOLD_FONT
    style_header_row(ws_sum, 4, HEADER_FILL_PRIMARY)

    for idx, (label, val, fmt) in enumerate(kpi_rows, start=5):
        c_lbl = ws_sum.cell(row=idx, column=1, value=label)
        c_lbl.font = NORMAL_FONT
        c_lbl.border = THIN_BORDER
        c_val = ws_sum.cell(row=idx, column=2, value=val)
        c_val.font = BOLD_FONT
        c_val.number_format = fmt
        c_val.border = THIN_BORDER
        
    auto_fit_columns(ws_sum)

    # -------------------------------------------------------------
    # 2. SHEET: Employee Comparison (Alphabetical)
    # -------------------------------------------------------------
    ws_emp = wb.create_sheet(title="Employee Comparison")
    ws_emp.views.sheetView[0].showGridLines = True
    
    headers = ["Employee Name", "Category", "Salary Head", "Current Month (₹)", "Samarth (₹)", "Difference (₹)", "Status"]
    for c_idx, h in enumerate(headers, start=1):
        ws_emp.cell(row=1, column=c_idx, value=h)
    style_header_row(ws_emp, 1, HEADER_FILL_PRIMARY)

    curr_row = 2
    for emp in grouped_emps:
        for hd in emp["heads"]:
            ws_emp.cell(row=curr_row, column=1, value=emp["employee_name"]).font = BOLD_FONT
            ws_emp.cell(row=curr_row, column=2, value=emp["category"]).font = NORMAL_FONT
            ws_emp.cell(row=curr_row, column=3, value=hd["salary_head"]).font = NORMAL_FONT
            
            c4 = ws_emp.cell(row=curr_row, column=4, value=hd["current_value"])
            c4.font = NORMAL_FONT
            c4.number_format = CURRENCY_FORMAT
            
            c5 = ws_emp.cell(row=curr_row, column=5, value=hd["samarth_value"])
            c5.font = NORMAL_FONT
            c5.number_format = CURRENCY_FORMAT
            
            c6 = ws_emp.cell(row=curr_row, column=6, value=hd["difference"])
            c6.font = BOLD_FONT
            c6.number_format = CURRENCY_FORMAT
            
            c7 = ws_emp.cell(row=curr_row, column=7, value=hd["status"])
            c7.font = BOLD_FONT
            if hd["status"] == "MATCH":
                c7.fill = FILL_MATCH
            elif hd["status"] == "MISMATCH":
                c7.fill = FILL_MISMATCH
            else:
                c7.fill = FILL_NOT_AVAIL
                
            for c in range(1, 8):
                ws_emp.cell(row=curr_row, column=c).border = THIN_BORDER
                if hd["is_total"]:
                    ws_emp.cell(row=curr_row, column=c).fill = FILL_TOTAL
                    
            curr_row += 1

    auto_fit_columns(ws_emp)

    # -------------------------------------------------------------
    # 3. SHEET: Salary Head Discrepancies
    # -------------------------------------------------------------
    ws_heads = wb.create_sheet(title="Salary Head Discrepancies")
    ws_heads.views.sheetView[0].showGridLines = True
    
    h_headers = ["Salary Head", "Category", "Employees with Mismatch", "Total Net Difference (₹)"]
    for c_idx, h in enumerate(h_headers, start=1):
        ws_heads.cell(row=1, column=c_idx, value=h)
    style_header_row(ws_heads, 1, HEADER_FILL_SECONDARY)

    for r_idx, hm in enumerate(head_mismatches, start=2):
        ws_heads.cell(row=r_idx, column=1, value=hm["salary_head"]).font = BOLD_FONT
        ws_heads.cell(row=r_idx, column=2, value=hm["category"]).font = NORMAL_FONT
        c3 = ws_heads.cell(row=r_idx, column=3, value=hm["employees_changed"])
        c3.font = BOLD_FONT
        c3.number_format = INTEGER_FORMAT
        c4 = ws_heads.cell(row=r_idx, column=4, value=hm["total_difference"])
        c4.font = BOLD_FONT
        c4.number_format = CURRENCY_FORMAT
        
        for c in range(1, 5):
            ws_heads.cell(row=r_idx, column=c).border = THIN_BORDER

    auto_fit_columns(ws_heads)

    # -------------------------------------------------------------
    # 5. SHEET: One-Time Entries (EXCLUDING Gross, Total Deductions, Net Pay)
    # -------------------------------------------------------------
    ws_ote = wb.create_sheet(title="One-Time Entries")
    ws_ote.views.sheetView[0].showGridLines = True
    
    ote_headers = ["S.No.", "Employee ID", "Employee Name", "Category", "Salary Head", "Entry Type", "Current Month (₹)", "Samarth Value (₹)", "Adjustment Amount (₹)", "Direction", "Remarks"]
    for c_idx, h in enumerate(ote_headers, start=1):
        ws_ote.cell(row=1, column=c_idx, value=h)
    style_header_row(ws_ote, 1, HEADER_FILL_SECONDARY)
    
    ote_rows = comparison_result.get("one_time_entries", [])
    for idx, ote in enumerate(ote_rows, start=2):
        ws_ote.cell(row=idx, column=1, value=ote["s_no"]).alignment = Alignment(horizontal="center")
        ws_ote.cell(row=idx, column=2, value=ote["employee_id"]).font = BOLD_FONT
        ws_ote.cell(row=idx, column=3, value=ote["employee_name"]).font = BOLD_FONT
        ws_ote.cell(row=idx, column=4, value=ote["category"]).font = NORMAL_FONT
        ws_ote.cell(row=idx, column=5, value=ote["salary_head"]).font = BOLD_FONT
        ws_ote.cell(row=idx, column=6, value=ote["entry_type"]).font = NORMAL_FONT
        
        c7 = ws_ote.cell(row=idx, column=7, value=ote["current_value"])
        c7.font = NORMAL_FONT
        c7.number_format = CURRENCY_FORMAT
        
        c8 = ws_ote.cell(row=idx, column=8, value=ote["samarth_value"])
        c8.font = NORMAL_FONT
        c8.number_format = CURRENCY_FORMAT
        
        c9 = ws_ote.cell(row=idx, column=9, value=ote["adjustment_amount"])
        c9.font = BOLD_FONT
        c9.number_format = CURRENCY_FORMAT
        
        ws_ote.cell(row=idx, column=10, value=ote["adjustment_direction"]).font = NORMAL_FONT
        ws_ote.cell(row=idx, column=11, value=ote["remarks"]).font = NORMAL_FONT
        
        for c in range(1, 12):
            ws_ote.cell(row=idx, column=c).border = THIN_BORDER
            
    auto_fit_columns(ws_ote)

    # Save to BytesIO
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out

def generate_one_time_entries_excel(comparison_result: Dict[str, Any]) -> io.BytesIO:
    """
    Generates a dedicated One-Time Payroll Entries Excel workbook for finance adjustments.
    EXCLUDES Gross / Total Earnings, Total Deductions, and Net Pay totals.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "One-Time Entries"
    ws.views.sheetView[0].showGridLines = True
    
    summary = comparison_result.get("summary", {})
    ote_rows = comparison_result.get("one_time_entries", [])
    
    # Title Block
    ws["A1"] = "ONE-TIME PAYROLL ENTRIES & ADJUSTMENTS REPORT"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated from Mismatches (Excludes Calculated Totals) | Files: {summary.get('current_filename', 'Current Month')} vs {summary.get('samarth_filename', 'Samarth')}"
    ws["A2"].font = SUBTITLE_FONT
    
    headers = ["S.No.", "Employee ID", "Employee Name", "Category", "Salary Head", "Entry Type", "Current Month (₹)", "Samarth Value (₹)", "Adjustment Amount (₹)", "Adjustment Direction", "Remarks"]
    for c_idx, h in enumerate(headers, start=1):
        ws.cell(row=4, column=c_idx, value=h)
    style_header_row(ws, 4, HEADER_FILL_SECONDARY)
    
    curr_row = 5
    total_adj_amount = 0.0
    
    for idx, ote in enumerate(ote_rows, start=1):
        ws.cell(row=curr_row, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=curr_row, column=2, value=ote.get("employee_id", "")).font = BOLD_FONT
        ws.cell(row=curr_row, column=3, value=ote.get("employee_name", "")).font = BOLD_FONT
        ws.cell(row=curr_row, column=4, value=ote.get("category", "")).font = NORMAL_FONT
        ws.cell(row=curr_row, column=5, value=ote.get("salary_head", "")).font = BOLD_FONT
        ws.cell(row=curr_row, column=6, value=ote.get("entry_type", "")).font = NORMAL_FONT
        
        c7 = ws.cell(row=curr_row, column=7, value=ote.get("current_value", 0.0))
        c7.font = NORMAL_FONT
        c7.number_format = CURRENCY_FORMAT
        
        c8 = ws.cell(row=curr_row, column=8, value=ote.get("samarth_value", 0.0))
        c8.font = NORMAL_FONT
        c8.number_format = CURRENCY_FORMAT
        
        adj_val = ote.get("adjustment_amount", 0.0)
        total_adj_amount += adj_val
        c9 = ws.cell(row=curr_row, column=9, value=adj_val)
        c9.font = BOLD_FONT
        c9.number_format = CURRENCY_FORMAT
        
        ws.cell(row=curr_row, column=10, value=ote.get("adjustment_direction", "")).font = NORMAL_FONT
        ws.cell(row=curr_row, column=11, value=ote.get("remarks", "")).font = NORMAL_FONT
        
        for c in range(1, 12):
            ws.cell(row=curr_row, column=c).border = THIN_BORDER
        curr_row += 1
        
    # Total Row
    if ote_rows:
        ws.cell(row=curr_row, column=1, value="").font = BOLD_FONT
        ws.cell(row=curr_row, column=2, value="TOTALS").font = BOLD_FONT
        ws.cell(row=curr_row, column=3, value=f"{len(ote_rows)} Adjustments").font = BOLD_FONT
        ws.cell(row=curr_row, column=4, value="").font = BOLD_FONT
        ws.cell(row=curr_row, column=5, value="").font = BOLD_FONT
        ws.cell(row=curr_row, column=6, value="").font = BOLD_FONT
        ws.cell(row=curr_row, column=7, value="").font = BOLD_FONT
        ws.cell(row=curr_row, column=8, value="").font = BOLD_FONT
        
        c_tot = ws.cell(row=curr_row, column=9, value=total_adj_amount)
        c_tot.font = BOLD_FONT
        c_tot.number_format = CURRENCY_FORMAT
        c_tot.fill = FILL_TOTAL
        
        ws.cell(row=curr_row, column=10, value="").font = BOLD_FONT
        ws.cell(row=curr_row, column=11, value="").font = BOLD_FONT
        
        for c in range(1, 12):
            ws.cell(row=curr_row, column=c).border = THIN_BORDER
            ws.cell(row=curr_row, column=c).fill = FILL_TOTAL
            
    auto_fit_columns(ws)
    
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
