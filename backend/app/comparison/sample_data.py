import io
from typing import Dict, Any
import openpyxl
from app.comparison.excel_parser import parse_excel_file
from app.comparison.comparator import compare_monthly_datasets, get_comparison_months

def generate_sample_month_workbook(month_idx: int) -> bytes:
    """
    Generates realistic Excel workbooks:
    - Month 0 (Month N-2, e.g. April): Category-wise historical sheets (Faulty Staff, Staff with duplicate NPS, Contractual Staff).
    - Month 1 (Month N-1, e.g. May): Category-wise historical sheets (increments & DA revisions).
    - Month 2 (Month N, e.g. June): Salary Generated (Samarth) Report structure.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    if month_idx < 2:
        # -------------------------------------------------------------
        # 1. FAULTY STAFF SHEET (Historical Category Format)
        # -------------------------------------------------------------
        ws_faulty = wb.create_sheet(title="Faulty Staff")
        faulty_headers = [
            "Sl. No.", "Particulars", "BASIC PAY", "DEARNESS ALLOWANCE", "HOUSE RENT ALLOWANCE",
            "TRANSPORT ALLOWANCE", "DA ON TPTA", "DEAN ALLOWANCE", "WARDEN ALLOWANCE",
            "ARREARS ON SALARY", "CHILDREN EDUCATION ALLOWANCE", "INCOME FROM CONSULTANCY PROJECT",
            "OTHER EARNINGS", "Total Earnings", "TAX DED. AT SOURCE", "NATIONAL PENSION SCHEME",
            "PROFESSIONAL TAX", "LICENSE FEE", "ELECTRICITY CHRGS", "WATER CHARGES",
            "IIT DH CLUB SUBSCRIPTION", "CCTI", "ELECTRICITY FIXED CHARGES", "MEDICAL RECOVERY",
            "RECOVERY OF OFFICE", "NPS RECOVERY", "Total Deductions", "Net Amount"
        ]
        ws_faulty.append(faulty_headers)
        
        faulty_rows = [
            (1, "Dr. Rajesh Sharma", 144200 + (10000 if month_idx >= 1 else 0), 0.46 + (0.04 if month_idx >= 1 else 0), 38934, 7200, 3500, 0, 25000, 18000, 200, 500),
            (2, "Dr. Ananya Iyer", 131400, 0.46 + (0.04 if month_idx >= 1 else 0), 35478, 7200, 0, 2500 if month_idx >= 1 else 0, 22000, 16500, 200, 500),
            (3, "Dr. Vikram Seth", 157600, 0.46 + (0.04 if month_idx >= 1 else 0), 42552, 7200, 4000, 0, 29000, 19800, 200, 500),
            (4, "Dr. Sneha Patil", 123800, 0.46 + (0.04 if month_idx >= 1 else 0), 33426, 7200, 0, 0, 19000, 15500, 200, 500),
        ]
        
        for item in faulty_rows:
            sl, name, basic, da_rate, hra, tpta, dean, warden, tds, nps, ptax, club = item
            da = round(basic * da_rate, 2)
            da_tpta = round(tpta * da_rate, 2)
            tot_earnings = basic + da + hra + tpta + da_tpta + dean + warden
            license_fee = 1190
            elec = 850 + (120 * month_idx)
            water = 250
            ccti = 300
            elec_fixed = 150
            tot_deductions = tds + nps + ptax + license_fee + elec + water + club + ccti + elec_fixed
            net_amount = tot_earnings - tot_deductions
            
            ws_faulty.append([
                sl, name, basic, da, hra, tpta, da_tpta, dean, warden,
                0, 0, 0, 0, tot_earnings,
                tds, nps, ptax, license_fee, elec, water, club, ccti, elec_fixed,
                0, 0, 0, tot_deductions, net_amount
            ])

        # -------------------------------------------------------------
        # 2. STAFF SHEET (with duplicate NATIONAL PENSION SCHEME)
        # -------------------------------------------------------------
        ws_staff = wb.create_sheet(title="Staff")
        staff_headers = [
            "Sl. No.", "Particulars", "BASIC PAY", "DEARNESS ALLOWANCE", "HOUSE RENT ALLOWANCE",
            "TRANSPORT ALLOWANCE", "DA ON TPTA", "ARREARS ON SALARY", "HIGHER EDUCATION INCENTIVE",
            "SALARY ARREARS HRA", "INCOME FROM CONSULTANCY PROJECT", "MOBILE CHARGE ALLOWANCE",
            "Total Earnings", "TAX DED. AT SOURCE", "NATIONAL PENSION SCHEME", "PROFESSIONAL TAX",
            "LICENSE FEE", "ELECTRICITY CHRGS", "WATER CHARGES", "CCTI", "ELECTRICITY FIXED CHARGES",
            "GPF", "NATIONAL PENSION SCHEME", "USAGE CHARGES", "Total Deductions", "Net Amount"
        ]
        ws_staff.append(staff_headers)
        
        staff_rows = [
            (1, "Mr. Ramesh Kumar", 65400 + (3000 if month_idx >= 1 else 0), 0.46 + (0.04 if month_idx >= 1 else 0), 17658, 3600, 1000, 500, 8000, 8175, 200, 3500),
            (2, "Ms. Pooja Hegde", 53600, 0.46 + (0.04 if month_idx >= 1 else 0), 14472, 3600, 0, 500, 5000, 6700, 200, 0),
            (3, "Mr. Santosh Nair", 78800 + (4000 if month_idx >= 1 else 0), 0.46 + (0.04 if month_idx >= 1 else 0), 21276, 7200, 2000, 500, 12000, 9850, 200, 5000),
        ]
        
        for item in staff_rows:
            sl, name, basic, da_rate, hra, tpta, edu_inc, mob, tds, nps_1, ptax, gpf = item
            da = round(basic * da_rate, 2)
            da_tpta = round(tpta * da_rate, 2)
            tot_earnings = basic + da + hra + tpta + da_tpta + edu_inc + mob
            nps_2 = round(nps_1 * 0.1, 2)
            tot_deductions = tds + nps_1 + nps_2 + ptax + 600 + 450 + 200 + 300 + 100 + gpf
            net_amount = tot_earnings - tot_deductions
            
            ws_staff.append([
                sl, name, basic, da, hra, tpta, da_tpta, 0, edu_inc,
                0, 0, mob, tot_earnings,
                tds, nps_1, ptax, 600, 450, 200, 300, 100,
                gpf, nps_2, 0, tot_deductions, net_amount
            ])

        # -------------------------------------------------------------
        # 3. CONTRACTUAL STAFF SHEET
        # -------------------------------------------------------------
        ws_contract = wb.create_sheet(title="Contractual Staff")
        contract_headers = [
            "Sl. No.", "Particulars", "BASIC PAY", "ARREARS ON SALARY", "Total Earnings",
            "TAX DED. AT SOURCE", "PROFESSIONAL TAX", "ACCOMMODATION CHARGES", "OTHER DEDUCTIONS",
            "Total Deductions", "Net Amount"
        ]
        ws_contract.append(contract_headers)
        
        contract_rows = [
            (1, "Mr. Ajay Mane", 35000 + (2000 if month_idx >= 1 else 0), 0, 1500, 200, 3000, 0),
            (2, "Ms. Deepa Deshmukh", 32000, 0, 1000, 200, 3000, 0),
        ]
        
        for item in contract_rows:
            sl, name, basic, arrears, tds, ptax, accomm, other_ded = item
            tot_earnings = basic + arrears
            tot_deductions = tds + ptax + accomm + other_ded
            net_amount = tot_earnings - tot_deductions
            ws_contract.append([
                sl, name, basic, arrears, tot_earnings,
                tds, ptax, accomm, other_ded, tot_deductions, net_amount
            ])

    else:
        # -------------------------------------------------------------
        # CURRENT MONTH: SALARY GENERATED (SAMARTH) REPORT STRUCTURE
        # -------------------------------------------------------------
        ws_samarth = wb.create_sheet(title="Salary Generated Report")
        samarth_headers = [
            "S_No.", "Id", "Name", "Category", "Department", "Designation",
            "Year", "Month", "Basic", "Dearness Allowance (DA)", "Transport Allowance (TA)",
            "House Rent Allowance (HRA)", "DA on TA", "Dean Allowance", "Warden Allowance",
            "Gross", "Income Tax", "National Pension Scheme (NPS)", "Professional Tax",
            "License Fee", "Electricity Charges", "Water Charge", "CCTI", "IIT DH Club Subscription",
            "Electricity Fixed Charges", "Deductions", "Net Pay",
            "NPS Employer Contribution", "CPF Employer Contribution", "EPF Employer Contribution"
        ]
        ws_samarth.append(samarth_headers)

        samarth_employees = [
            # Faulty Staff
            (1, "EMP101", "Dr. Rajesh Sharma", "Faulty Staff", "CSE", "Professor", 154200, 77100, 7200, 38934, 3600, 3500, 0, 26000, 20000, 200, 1190, 1090, 250, 300, 500, 150, 21588, 0, 0),
            (2, "EMP102", "Dr. Ananya Iyer", "Faulty Staff", "EE", "Associate Professor", 131400, 65700, 7200, 35478, 3600, 0, 2500, 23000, 18000, 200, 1190, 970, 250, 300, 500, 150, 18396, 0, 0),
            (3, "EMP103", "Dr. Vikram Seth", "Faulty Staff", "ME", "Professor", 162600, 81300, 7200, 42552, 3600, 4000, 0, 31000, 21000, 200, 1190, 1090, 250, 300, 500, 150, 22764, 0, 0),
            (4, "EMP104", "Dr. Sneha Patil", "Faulty Staff", "CH", "Assistant Professor", 128800, 64400, 7200, 33426, 3600, 0, 0, 20000, 16500, 200, 1190, 970, 250, 300, 500, 150, 18032, 0, 0),
            # Staff
            (5, "EMP201", "Mr. Ramesh Kumar", "Staff", "Admin", "Assistant Registrar", 68400, 34200, 3600, 17658, 1800, 0, 0, 8500, 9576, 200, 600, 450, 200, 300, 0, 100, 9576, 0, 0),
            (6, "EMP202", "Ms. Pooja Hegde", "Staff", "Accounts", "Junior Assistant", 53600, 26800, 3600, 14472, 1800, 0, 0, 5200, 7504, 200, 600, 450, 200, 300, 0, 100, 7504, 0, 0),
            (7, "EMP203", "Mr. Santosh Nair", "Staff", "Estate", "Superintendent Engineer", 82800, 41400, 7200, 21276, 3600, 0, 0, 13000, 11592, 200, 600, 450, 200, 300, 0, 100, 11592, 0, 0),
            # Contractual Staff
            (8, "EMP301", "Mr. Ajay Mane", "Contractual Staff", "Security", "Security Supervisor", 37000, 0, 0, 0, 0, 0, 0, 1600, 0, 200, 0, 0, 0, 0, 0, 0, 0, 0, 4440),
            (9, "EMP302", "Ms. Deepa Deshmukh", "Contractual Staff", "Library", "Library Trainee", 32000, 0, 0, 0, 0, 0, 0, 1000, 0, 200, 0, 0, 0, 0, 0, 0, 0, 0, 3840),
        ]

        for emp in samarth_employees:
            s_no, emp_id, name, cat, dept, desig, basic, da, ta, hra, da_ta, dean, warden, tds, nps, ptax, lic, elec, water, ccti, club, elec_fixed, nps_emp, cpf_emp, epf_emp = emp
            gross = basic + da + ta + hra + da_ta + dean + warden
            deductions = tds + nps + ptax + lic + elec + water + ccti + club + elec_fixed
            net_pay = gross - deductions
            ws_samarth.append([
                s_no, emp_id, name, cat, dept, desig,
                2026, 6, basic, da, ta, hra, da_ta, dean, warden,
                gross, tds, nps, ptax, lic, elec, water, ccti, club, elec_fixed,
                deductions, net_pay, nps_emp, cpf_emp, epf_emp
            ])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def get_sample_comparison_dataset(current_year: int = 2026, current_month: int = 6) -> Dict[str, Any]:
    """Builds and returns complete 3-month sample comparison result directly based on dynamic processing month."""
    months_meta = get_comparison_months(current_year, current_month, count=3)
    monthly_datasets = []
    
    for idx, m_info in enumerate(months_meta):
        label = m_info["label"]
        file_bytes = generate_sample_month_workbook(idx)
        cat_data = parse_excel_file(file_bytes, f"Salary_{label.replace(' ', '_')}.xlsx")
        monthly_datasets.append({
            "month_index": idx,
            "month_label": label,
            "categories": cat_data
        })
        
    return compare_monthly_datasets(monthly_datasets)
