import io
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Styles definition
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Arial", size=14, bold=True, color="1E293B")
SUBTITLE_FONT = Font(name="Arial", size=10, italic=True, color="64748B")
BOLD_FONT = Font(name="Arial", size=10, bold=True, color="0F172A")
NORMAL_FONT = Font(name="Arial", size=10, color="0F172A")

HEADER_FILL_PRIMARY = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid") # Deep blue
HEADER_FILL_SECONDARY = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid") # Sky blue
HEADER_FILL_ACCENT = PatternFill(start_color="475569", end_color="475569", fill_type="solid") # Slate

FILL_INCREASED = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Soft green
FILL_DECREASED = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Soft red
FILL_TOTAL_ROW = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid") # Soft gray

THIN_BORDER = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

CURRENCY_FORMAT = '"₹"#,##0.00;[Red]("-₹"#,##0.00);"-"'
INTEGER_FORMAT = '#,##0'

def auto_fit_columns(ws, min_width=12, max_width=45):
    """Auto-fit column widths with safety boundaries."""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if '\n' in val_str:
                val_str = max(val_str.split('\n'), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 4, max_width))

def style_header_row(ws, row_idx: int, fill=HEADER_FILL_PRIMARY):
    for cell in ws[row_idx]:
        cell.font = HEADER_FONT
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

def generate_comparison_excel_report(comparison_data: Dict[str, Any]) -> io.BytesIO:
    """
    Generates a rich, multi-sheet comparison Excel report from comparison dataset.
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    months = comparison_data.get("months", [])
    num_months = len(months)
    overall_summary = comparison_data.get("overall_summary", {})
    transitions = comparison_data.get("transitions_summary", [])
    head_changes = comparison_data.get("head_changes_summary", [])
    employee_changes = comparison_data.get("employee_changes", [])
    presence_changes = comparison_data.get("presence_changes", [])
    
    # -------------------------------------------------------------
    # 1. SHEET: Summary & Overview
    # -------------------------------------------------------------
    ws_sum = wb.create_sheet(title="Summary & Overview")
    ws_sum.views.sheetView[0].showGridLines = True
    
    ws_sum["A1"] = "MONTHLY SALARY EXCEL COMPARISON REPORT"
    ws_sum["A1"].font = TITLE_FONT
    ws_sum["A2"] = f"Comparison Periods: {' → '.join(months)}"
    ws_sum["A2"].font = SUBTITLE_FONT
    
    # KPI Table
    ws_sum["A4"] = "Overall Metric"
    ws_sum["B4"] = "Value"
    style_header_row(ws_sum, 4, HEADER_FILL_PRIMARY)
    
    kpis = [
        ("Total Employees Analyzed", overall_summary.get("total_employees", 0)),
        ("Faulty Staff Count", overall_summary.get("faulty_staff_count", 0)),
        ("Staff Count", overall_summary.get("staff_count", 0)),
        ("Contractual Staff Count", overall_summary.get("contractual_staff_count", 0)),
        ("Total Salary Heads Compared", overall_summary.get("total_salary_heads_compared", 0)),
        ("Total Value Changes Detected", overall_summary.get("total_value_changes", 0)),
        ("Total Increases (+)", overall_summary.get("total_increases", 0)),
        ("Total Decreases (-)", overall_summary.get("total_decreases", 0)),
        ("Total No-Change Instances", overall_summary.get("total_no_change", 0)),
        ("New / Missing Employees", overall_summary.get("total_new_or_missing_employees", 0)),
    ]
    
    for r_idx, (label, val) in enumerate(kpis, start=5):
        ws_sum.cell(row=r_idx, column=1, value=label).font = NORMAL_FONT
        cell_v = ws_sum.cell(row=r_idx, column=2, value=val)
        cell_v.font = BOLD_FONT
        cell_v.number_format = INTEGER_FORMAT
        cell_v.alignment = Alignment(horizontal="right")
        ws_sum.cell(row=r_idx, column=1).border = THIN_BORDER
        cell_v.border = THIN_BORDER
        
    # Month Transition Summary Table
    start_r = 17
    ws_sum.cell(row=start_r, column=1, value="Month Transition Overview").font = TITLE_FONT
    start_r += 1
    
    t_headers = ["Period Transition", "Changed Values", "Increased Values", "Decreased Values", "Employees Affected"]
    for c_idx, th in enumerate(t_headers, start=1):
        ws_sum.cell(row=start_r, column=c_idx, value=th)
    style_header_row(ws_sum, start_r, HEADER_FILL_SECONDARY)
    
    for t_idx, tr in enumerate(transitions, start=start_r + 1):
        ws_sum.cell(row=t_idx, column=1, value=tr["transition_label"]).font = BOLD_FONT
        ws_sum.cell(row=t_idx, column=2, value=tr["changed_values"]).number_format = INTEGER_FORMAT
        ws_sum.cell(row=t_idx, column=3, value=tr["increased_values"]).number_format = INTEGER_FORMAT
        ws_sum.cell(row=t_idx, column=4, value=tr["decreased_values"]).number_format = INTEGER_FORMAT
        ws_sum.cell(row=t_idx, column=5, value=tr["affected_employees_count"]).number_format = INTEGER_FORMAT
        for c in range(1, 6):
            ws_sum.cell(row=t_idx, column=c).border = THIN_BORDER
            ws_sum.cell(row=t_idx, column=c).font = NORMAL_FONT if c > 1 else BOLD_FONT
            
    auto_fit_columns(ws_sum)
    
    # -------------------------------------------------------------
    # 2. SHEET: Salary Head Changes
    # -------------------------------------------------------------
    ws_heads = wb.create_sheet(title="Salary Head Changes")
    ws_heads.views.sheetView[0].showGridLines = True
    
    head_cols = ["Category", "Salary Head", "Type"]
    for tr in transitions:
        head_cols.append(f"{tr['transition_label']} (Changed Count)")
    head_cols.extend(["Overall Changed Count", "Increased Count", "Decreased Count", "Net Sum Difference (₹)"])
    
    for c_idx, h in enumerate(head_cols, start=1):
        ws_heads.cell(row=1, column=c_idx, value=h)
    style_header_row(ws_heads, 1, HEADER_FILL_PRIMARY)
    
    for r_idx, hd in enumerate(head_changes, start=2):
        col_c = 1
        ws_heads.cell(row=r_idx, column=col_c, value=hd["category"]).font = NORMAL_FONT; col_c += 1
        ws_heads.cell(row=r_idx, column=col_c, value=hd["salary_head"]).font = BOLD_FONT if hd.get("is_total") else NORMAL_FONT; col_c += 1
        type_str = "Total / Aggregate" if hd.get("is_total") else "Employer Contribution" if hd.get("is_employer_contribution") else "Standard Salary Head"
        ws_heads.cell(row=r_idx, column=col_c, value=type_str).font = NORMAL_FONT; col_c += 1
        
        # Transition changed counts
        for tc in hd.get("transition_changes_count", []):
            c = ws_heads.cell(row=r_idx, column=col_c, value=tc)
            c.number_format = INTEGER_FORMAT
            c.font = BOLD_FONT if tc > 0 else NORMAL_FONT
            col_c += 1
            
        # Overall stats
        c1 = ws_heads.cell(row=r_idx, column=col_c, value=hd.get("overall_changes_count", 0)); c1.number_format = INTEGER_FORMAT; col_c += 1
        c2 = ws_heads.cell(row=r_idx, column=col_c, value=hd.get("increased_count", 0)); c2.number_format = INTEGER_FORMAT; col_c += 1
        c3 = ws_heads.cell(row=r_idx, column=col_c, value=hd.get("decreased_count", 0)); c3.number_format = INTEGER_FORMAT; col_c += 1
        c4 = ws_heads.cell(row=r_idx, column=col_c, value=hd.get("net_diff", 0.0)); c4.number_format = CURRENCY_FORMAT; col_c += 1
        
        for c in range(1, col_c):
            ws_heads.cell(row=r_idx, column=c).border = THIN_BORDER
            if hd.get("is_total"):
                ws_heads.cell(row=r_idx, column=c).fill = FILL_TOTAL_ROW
                
    auto_fit_columns(ws_heads)
    
    # -------------------------------------------------------------
    # Helper to populate detailed employee sheets
    # -------------------------------------------------------------
    def populate_employee_sheet(ws_target, rows_data):
        ws_target.views.sheetView[0].showGridLines = True
        
        emp_cols = ["Category", "Employee Name", "Salary Head", "Type"]
        for m in months:
            emp_cols.append(m)
        for tr in transitions:
            emp_cols.append(f"Diff ({tr['transition_label']})")
            emp_cols.append(f"Status ({tr['transition_label']})")
        emp_cols.extend(["Overall Diff (₹)", "Overall Status"])
        
        for c_idx, h in enumerate(emp_cols, start=1):
            ws_target.cell(row=1, column=c_idx, value=h)
        style_header_row(ws_target, 1, HEADER_FILL_PRIMARY)
        
        for r_idx, row in enumerate(rows_data, start=2):
            col_c = 1
            ws_target.cell(row=r_idx, column=col_c, value=row["category"]).font = NORMAL_FONT; col_c += 1
            ws_target.cell(row=r_idx, column=col_c, value=row["employee_name"]).font = BOLD_FONT; col_c += 1
            ws_target.cell(row=r_idx, column=col_c, value=row["salary_head"]).font = NORMAL_FONT; col_c += 1
            ws_target.cell(row=r_idx, column=col_c, value="Total" if row["is_total"] else "Head").font = NORMAL_FONT; col_c += 1
            
            # Monthly values
            for mv in row["month_values"]:
                c = ws_target.cell(row=r_idx, column=col_c, value=mv)
                c.number_format = CURRENCY_FORMAT
                c.font = NORMAL_FONT
                col_c += 1
                
            # Transition diffs
            for tr_d in row["transitions"]:
                c_diff = ws_target.cell(row=r_idx, column=col_c, value=tr_d["diff"])
                c_diff.number_format = CURRENCY_FORMAT
                c_diff.font = BOLD_FONT if tr_d["status"] != "NO_CHANGE" else NORMAL_FONT
                col_c += 1
                
                c_st = ws_target.cell(row=r_idx, column=col_c, value=tr_d["status"])
                c_st.font = NORMAL_FONT
                if tr_d["status"] == "INCREASED":
                    c_st.fill = FILL_INCREASED
                elif tr_d["status"] == "DECREASED":
                    c_st.fill = FILL_DECREASED
                col_c += 1
                
            # Overall diff and status
            c_ov_diff = ws_target.cell(row=r_idx, column=col_c, value=row["overall_diff"])
            c_ov_diff.number_format = CURRENCY_FORMAT
            c_ov_diff.font = BOLD_FONT if row["overall_status"] != "NO_CHANGE" else NORMAL_FONT
            col_c += 1
            
            c_ov_st = ws_target.cell(row=r_idx, column=col_c, value=row["overall_status"])
            c_ov_st.font = NORMAL_FONT
            if row["overall_status"] == "INCREASED":
                c_ov_st.fill = FILL_INCREASED
            elif row["overall_status"] == "DECREASED":
                c_ov_st.fill = FILL_DECREASED
            col_c += 1
            
            for c in range(1, col_c):
                ws_target.cell(row=r_idx, column=c).border = THIN_BORDER
                if row["is_total"]:
                    ws_target.cell(row=r_idx, column=c).fill = FILL_TOTAL_ROW
                    
        auto_fit_columns(ws_target)

    # 3. SHEET: All Changes (Changed Only)
    changed_rows = [r for r in employee_changes if r.get("has_change", False)]
    ws_all_changed = wb.create_sheet(title="All Value Changes")
    populate_employee_sheet(ws_all_changed, changed_rows)
    
    # 4. SHEET: Faulty Staff Changes
    faulty_rows = [r for r in employee_changes if r.get("category_key") == "faulty_staff" and r.get("has_change", False)]
    ws_faulty = wb.create_sheet(title="Faulty Staff Changes")
    populate_employee_sheet(ws_faulty, faulty_rows)
    
    # 5. SHEET: Staff Changes
    staff_rows = [r for r in employee_changes if r.get("category_key") == "staff" and r.get("has_change", False)]
    ws_staff = wb.create_sheet(title="Staff Changes")
    populate_employee_sheet(ws_staff, staff_rows)
    
    # 6. SHEET: Contractual Staff Changes
    contract_rows = [r for r in employee_changes if r.get("category_key") == "contractual_staff" and r.get("has_change", False)]
    ws_contract = wb.create_sheet(title="Contractual Staff Changes")
    populate_employee_sheet(ws_contract, contract_rows)
    
    # 7. SHEET: Increased Values Only
    inc_rows = [r for r in employee_changes if r.get("overall_status") == "INCREASED" and not r.get("is_total", False)]
    ws_inc = wb.create_sheet(title="Increased Values")
    populate_employee_sheet(ws_inc, inc_rows)
    
    # 8. SHEET: Decreased Values Only
    dec_rows = [r for r in employee_changes if r.get("overall_status") == "DECREASED" and not r.get("is_total", False)]
    ws_dec = wb.create_sheet(title="Decreased Values")
    populate_employee_sheet(ws_dec, dec_rows)
    
    # 9. SHEET: New & Missing Employees
    ws_pres = wb.create_sheet(title="New & Missing Employees")
    ws_pres.views.sheetView[0].showGridLines = True
    
    pres_cols = ["Category", "Employee Name", "Status / Movement", "Present In Months", "Missing In Months"]
    for c_idx, h in enumerate(pres_cols, start=1):
        ws_pres.cell(row=1, column=c_idx, value=h)
    style_header_row(ws_pres, 1, HEADER_FILL_PRIMARY)
    
    for r_idx, pr in enumerate(presence_changes, start=2):
        ws_pres.cell(row=r_idx, column=1, value=pr["category"]).font = NORMAL_FONT
        ws_pres.cell(row=r_idx, column=2, value=pr["employee_name"]).font = BOLD_FONT
        c_st = ws_pres.cell(row=r_idx, column=3, value=pr["status_label"])
        c_st.font = BOLD_FONT
        if "New" in pr["status_label"]:
            c_st.fill = FILL_INCREASED
        elif "Missing" in pr["status_label"] or "Exited" in pr["status_label"]:
            c_st.fill = FILL_DECREASED
            
        ws_pres.cell(row=r_idx, column=4, value=", ".join(pr["present_months"])).font = NORMAL_FONT
        ws_pres.cell(row=r_idx, column=5, value=", ".join(pr["missing_months"])).font = NORMAL_FONT
        for c in range(1, 6):
            ws_pres.cell(row=r_idx, column=c).border = THIN_BORDER
            
    auto_fit_columns(ws_pres)

    # 10. SHEET: Reconciliation
    reconciliation_rows = comparison_data.get("reconciliation", [])
    if reconciliation_rows:
        ws_recon = wb.create_sheet(title="Reconciliation")
        ws_recon.views.sheetView[0].showGridLines = True
        recon_cols = ["Employee Name", "Category", "Salary Head", "Historical Value (₹)", "Salary Generated Value (₹)", "Difference (₹)", "Reconciliation Status"]
        for c_idx, h in enumerate(recon_cols, start=1):
            ws_recon.cell(row=1, column=c_idx, value=h)
        style_header_row(ws_recon, 1, HEADER_FILL_PRIMARY)
        for r_idx, rc in enumerate(reconciliation_rows, start=2):
            ws_recon.cell(row=r_idx, column=1, value=rc["employee_name"]).font = BOLD_FONT
            ws_recon.cell(row=r_idx, column=2, value=rc["category"]).font = NORMAL_FONT
            ws_recon.cell(row=r_idx, column=3, value=rc["salary_head"]).font = NORMAL_FONT
            c_h = ws_recon.cell(row=r_idx, column=4, value=rc["historical_value"])
            c_h.font = NORMAL_FONT
            c_h.number_format = CURRENCY_FORMAT
            c_g = ws_recon.cell(row=r_idx, column=5, value=rc["salary_generated_value"])
            c_g.font = NORMAL_FONT
            c_g.number_format = CURRENCY_FORMAT
            c_d = ws_recon.cell(row=r_idx, column=6, value=rc["difference"])
            c_d.font = BOLD_FONT
            c_d.number_format = CURRENCY_FORMAT
            c_s = ws_recon.cell(row=r_idx, column=7, value=rc["status"])
            c_s.font = BOLD_FONT
            if rc["status"] == "MATCH":
                c_s.fill = FILL_INCREASED
            else:
                c_s.fill = FILL_DECREASED
            for c in range(1, 8):
                ws_recon.cell(row=r_idx, column=c).border = THIN_BORDER
        auto_fit_columns(ws_recon)

    # 11. SHEET: Field Mappings
    field_mappings = comparison_data.get("field_mappings", [])
    if field_mappings:
        ws_map = wb.create_sheet(title="Field Mappings")
        ws_map.views.sheetView[0].showGridLines = True
        map_cols = ["Historical Salary Head", "Salary Generated Field", "Classification", "Mapping Status"]
        for c_idx, h in enumerate(map_cols, start=1):
            ws_map.cell(row=1, column=c_idx, value=h)
        style_header_row(ws_map, 1, HEADER_FILL_ACCENT)
        for r_idx, fm in enumerate(field_mappings, start=2):
            ws_map.cell(row=r_idx, column=1, value=fm["historical_head"]).font = BOLD_FONT
            ws_map.cell(row=r_idx, column=2, value=fm["salary_generated_field"]).font = NORMAL_FONT
            classification = "Total / Structural" if fm["is_total"] else "Employer Contribution" if fm["is_employer_contribution"] else "Standard Salary Head"
            ws_map.cell(row=r_idx, column=3, value=classification).font = NORMAL_FONT
            ws_map.cell(row=r_idx, column=4, value=fm["status"]).font = BOLD_FONT
            for c in range(1, 5):
                ws_map.cell(row=r_idx, column=c).border = THIN_BORDER
        auto_fit_columns(ws_map)
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
